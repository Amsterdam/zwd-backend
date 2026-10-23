from datetime import date
from io import BytesIO
from apps.cases.admin import export_queryset_to_excel
from apps.cases.models import Case, CaseClose, CaseCloseReason
from apps.cases.retention import (
    filter_retention_expired,
    filter_retention_running,
    filter_successful_expired,
    filter_unknown_expired,
    filter_unsuccessful_expired,
    get_result_label,
    get_retention_expiry_date,
    get_retention_years,
)
from apps.homeownerassociation.models import HomeownerAssociation
from django.contrib.auth import get_user_model
from django.core import management
from django.test import TestCase
from django.utils import timezone
from model_bakery import baker
from openpyxl import load_workbook


def make_closed_case(is_successful, end_date, **kwargs):
    case = baker.make(Case, **kwargs)
    reason = baker.make(CaseCloseReason, is_successful=is_successful)
    baker.make(CaseClose, case=case, reason=reason)
    # CaseClose.save() sets end_date to today, so override it afterwards
    Case.objects.filter(id=case.id).update(end_date=end_date)
    return Case.objects.get(id=case.id)


class RetentionTest(TestCase):
    def setUp(self):
        management.call_command("flush", verbosity=0, interactive=False)
        super().setUp()

    def test_successful_case_has_five_year_retention(self):
        case = make_closed_case(True, date(2026, 10, 21))
        self.assertEqual(get_retention_years(case), 5)
        self.assertEqual(get_retention_expiry_date(case), date(2031, 12, 31))

    def test_unsuccessful_case_has_one_year_retention(self):
        case = make_closed_case(False, date(2026, 10, 21))
        self.assertEqual(get_retention_years(case), 1)
        self.assertEqual(get_retention_expiry_date(case), date(2027, 12, 31))

    def test_closed_case_without_case_close_has_five_year_retention(self):
        case = baker.make(Case, end_date=date(2020, 3, 1))
        self.assertEqual(get_retention_years(case), 5)
        self.assertEqual(get_retention_expiry_date(case), date(2025, 12, 31))
        self.assertEqual(get_result_label(case), "Onbekend")

    def test_open_case_has_no_retention(self):
        case = baker.make(Case, end_date=None)
        self.assertIsNone(get_retention_years(case))
        self.assertIsNone(get_retention_expiry_date(case))
        self.assertEqual(get_result_label(case), "")


class RetentionFilterTest(TestCase):
    def setUp(self):
        management.call_command("flush", verbosity=0, interactive=False)
        super().setUp()
        self.successful = make_closed_case(True, date(2022, 6, 1))
        self.unsuccessful = make_closed_case(False, date(2026, 10, 21))
        self.unknown = baker.make(Case, end_date=date(2022, 6, 1))
        self.open = baker.make(Case, end_date=None)
        self.archived = make_closed_case(False, date(2020, 1, 1))
        Case.objects.filter(id=self.archived.id).update(archived_at=timezone.now())

    def _ids(self, queryset):
        return set(queryset.values_list("id", flat=True))

    def test_nothing_expired_on_last_day_of_retention(self):
        today = date(2027, 12, 31)
        self.assertEqual(
            self._ids(filter_retention_expired(Case.objects.all(), today)), set()
        )
        self.assertEqual(
            self._ids(filter_retention_running(Case.objects.all(), today)),
            {self.successful.id, self.unsuccessful.id, self.unknown.id},
        )

    def test_expired_on_first_day_after_retention(self):
        today = date(2028, 1, 1)
        cases = Case.objects.all()
        self.assertEqual(
            self._ids(filter_successful_expired(cases, today)), {self.successful.id}
        )
        self.assertEqual(
            self._ids(filter_unsuccessful_expired(cases, today)), {self.unsuccessful.id}
        )
        self.assertEqual(
            self._ids(filter_unknown_expired(cases, today)), {self.unknown.id}
        )
        self.assertEqual(
            self._ids(filter_retention_expired(cases, today)),
            {self.successful.id, self.unsuccessful.id, self.unknown.id},
        )
        self.assertEqual(self._ids(filter_retention_running(cases, today)), set())

    def test_open_and_archived_cases_are_never_included(self):
        today = date(2100, 1, 1)
        cases = Case.objects.all()
        for ids in (
            self._ids(filter_retention_expired(cases, today)),
            self._ids(filter_retention_running(cases, today)),
        ):
            self.assertNotIn(self.open.id, ids)
            self.assertNotIn(self.archived.id, ids)

    def test_filters_match_retention_expiry_date(self):
        cases = Case.objects.exclude(id__in=[self.open.id, self.archived.id])
        for today in (date(2027, 12, 31), date(2028, 1, 1), date(2032, 6, 1)):
            expired_ids = self._ids(filter_retention_expired(cases, today))
            for case in cases:
                self.assertEqual(
                    case.id in expired_ids,
                    get_retention_expiry_date(case) < today,
                    f"case {case.id} on {today}",
                )

    def test_admin_changelist_filter(self):
        user = get_user_model().objects.create_superuser(
            username="admin", email="admin@example.com", password="admin"
        )
        self.client.force_login(user)
        response = self.client.get("/admin/cases/case/", {"bewaartermijn": "verlopen"})
        self.assertEqual(response.status_code, 200)
        expected = self._ids(filter_retention_expired(Case.objects.all()))
        self.assertEqual(
            {case.id for case in response.context["cl"].result_list}, expected
        )


class ExportQuerysetToExcelTest(TestCase):
    def setUp(self):
        management.call_command("flush", verbosity=0, interactive=False)
        super().setUp()

    def _export_rows(self, queryset):
        response = export_queryset_to_excel(None, None, queryset)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content))
        return list(workbook.active.iter_rows(values_only=True))

    def test_export_headers(self):
        rows = self._export_rows(Case.objects.none())
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], "Zaak ID")
        self.assertEqual(rows[0][-1], "Gearchiveerd op")

    def test_export_closed_case(self):
        homeowner_association = baker.make(
            HomeownerAssociation, name="VvE Teststraat", number_of_apartments=13
        )
        case = make_closed_case(
            False,
            date(2025, 4, 30),
            homeowner_association=homeowner_association,
            request_date=date(2025, 4, 8),
        )

        rows = self._export_rows(Case.objects.all())
        self.assertEqual(len(rows), 2)
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row["Zaak ID"], case.id)
        self.assertEqual(row["VvE"], "VvE Teststraat")
        self.assertEqual(row["Status"], "Afgesloten")
        self.assertEqual(row["Aanvraagdatum"], "08-04-2025")
        self.assertEqual(row["Einddatum"], "30-04-2025")
        self.assertEqual(row["Resultaat"], case.caseclose_set.get().reason.name)
        self.assertEqual(row["Succesvol"], "Nee")
        self.assertEqual(row["Bewaartermijn (jaren)"], 1)
        self.assertEqual(row["Bewaartermijn verloopt op"], "31-12-2026")
        self.assertIsNone(row["Gearchiveerd op"])

    def test_export_archived_case(self):
        case = make_closed_case(True, date(2020, 1, 1))
        archived_at = timezone.make_aware(timezone.datetime(2026, 1, 5, 10, 0))
        Case.objects.filter(id=case.id).update(archived_at=archived_at)

        rows = self._export_rows(Case.objects.all())
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row["Gearchiveerd op"], "05-01-2026")

    def test_export_closed_case_without_case_close(self):
        baker.make(Case, legacy_id="L-1", end_date=date(2019, 6, 1))

        rows = self._export_rows(Case.objects.all())
        row = dict(zip(rows[0], rows[1]))
        self.assertEqual(row["Resultaat"], "Onbekend")
        self.assertIsNone(row["Succesvol"])
        self.assertEqual(row["Bewaartermijn (jaren)"], 5)
        self.assertEqual(row["Bewaartermijn verloopt op"], "31-12-2024")

    def test_export_open_case_has_empty_retention(self):
        baker.make(Case)

        rows = self._export_rows(Case.objects.all())
        row = dict(zip(rows[0], rows[1]))
        self.assertIsNone(row["Einddatum"])
        self.assertIsNone(row["Resultaat"])
        self.assertIsNone(row["Bewaartermijn (jaren)"])
        self.assertIsNone(row["Bewaartermijn verloopt op"])
